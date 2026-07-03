from datetime import date as dt_date, timedelta
from types import SimpleNamespace
import calendar
from decimal import Decimal
from typing import Dict, List, Tuple, Set

from django.contrib.auth.decorators import login_required
from django.shortcuts import render
from django.http import HttpResponse
from django.contrib import messages
from django.db.models import Q

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
except Exception:
    openpyxl = None

from apps.organization.models import OrgUnit
from apps.attendance.models_batch import AttendanceCommit, AttendanceCommitItem
from apps.hr.models import Employee
from apps.backoffice.services.access_scope import get_allowed_attendance_units, unit_in_attendance_scope
try:
    from apps.hr.models.temp_assignment import TempAssignment
except Exception:
    TempAssignment = None


FULL_LEAVE_CODES = {"P", "O", "C", "B", "R"}
HALF_LEAVE_CODES = {"LP", "PL", "LB", "BL"}
SUPPLEMENT_OUT_CODE = "BS"


def _effective_temp_assignment_qs_between(start: dt_date, end: dt_date):
    """Điều động có hiệu lực trong khoảng tháng, dùng làm nguồn nghiệp vụ gốc cho BS."""
    if TempAssignment is None:
        return None
    return (
        TempAssignment.objects.filter(
            apply_flag=True,
            status__in=TempAssignment.effective_statuses(),
            start_date__lte=end,
        )
        .filter(Q(end_date__isnull=True) | Q(end_date__gte=start))
        .select_related("employee", "employee__team", "from_unit", "to_unit")
    )


def _iter_effective_dates_for_assignment(ta, start: dt_date, end: dt_date):
    cur = max(start, ta.start_date)
    last = min(end, ta.end_date or end)
    while cur <= last:
        yield cur
        cur += timedelta(days=1)


def _units_for_attendance(request):
    """Lấy danh sách đơn vị chấm công theo AccessControl."""
    return get_allowed_attendance_units(request.user)


def _month_range(year: int, month: int) -> Tuple[dt_date, dt_date, int]:
    """Trả về (start_date, end_date, days_in_month)."""
    days_in_month = calendar.monthrange(year, month)[1]
    start = dt_date(year, month, 1)
    end = dt_date(year, month, days_in_month)
    return start, end, days_in_month


def _parse_month_year(request) -> Tuple[int, int]:
    """Parse tháng/năm an toàn cho cả màn hình và Excel export."""
    from django.utils import timezone

    today = timezone.localdate()
    try:
        month = int(request.GET.get("month") or today.month)
    except Exception:
        month = today.month
    try:
        year = int(request.GET.get("year") or today.year)
    except Exception:
        year = today.year

    if month < 1 or month > 12:
        month = today.month
    if year < 2000 or year > 2100:
        year = today.year
    return month, year


def _weekday_label_precise(d: dt_date) -> str:
    wd = d.weekday()
    return ["T2", "T3", "T4", "T5", "T6", "T7", "CN"][wd]


def _to_decimal(value) -> Decimal:
    if value is None:
        return Decimal("0.00")
    if isinstance(value, Decimal):
        return value
    try:
        return Decimal(str(value))
    except Exception:
        return Decimal("0.00")


def _credit_from_item(it: AttendanceCommitItem | None, field_name: str, code_field_name: str) -> Decimal:
    """
    Ưu tiên snapshot trên dòng chốt; fallback về AttendanceCode để dữ liệu cũ vẫn xem được.

    Lưu ý: Nếu server đã tự sinh migration Phase 5, các dòng chốt cũ có thể có
    snapshot số = 0 nhưng code_snapshot còn trống. Khi đó phải fallback về AttendanceCode.
    """
    if not it:
        return Decimal("0.00")
    if getattr(it, "code_snapshot", ""):
        return _to_decimal(getattr(it, field_name, 0))
    code = getattr(it, "code", None)
    return _to_decimal(getattr(code, code_field_name, 0) if code else 0)


def _commit_item_code_text(it: AttendanceCommitItem | None) -> str:
    """Báo cáo tháng ưu tiên mã snapshot để dữ liệu lịch sử không đổi theo cấu hình hiện tại."""
    if not it:
        return ""
    return getattr(it, "code_snapshot", "") or (it.code.code if (it.code_id and it.code) else "")


def _is_supplement_in(it: AttendanceCommitItem | None) -> bool:
    return bool(it and str(getattr(it, "bs_direction", "")) == str(AttendanceCommitItem.BSDirection.IN))


def _is_supplement_out(it: AttendanceCommitItem | None) -> bool:
    return bool(it and str(getattr(it, "bs_direction", "")) == str(AttendanceCommitItem.BSDirection.OUT))


def _leave_credit(it: AttendanceCommitItem | None) -> Decimal:
    """
    Tính công nghỉ để thống kê nhanh:
    - P/O/C/B/R: 1 công nghỉ.
    - LP/PL/LB/BL: 0.5 công nghỉ.
    - BS đi/BS đến không tính là nghỉ.
    """
    if not it or _is_supplement_out(it) or _is_supplement_in(it) or not bool(getattr(it, "include_in_unit", True)):
        return Decimal("0.00")
    code_text = _commit_item_code_text(it).upper()
    if code_text in FULL_LEAVE_CODES:
        return Decimal("1.00")
    if code_text in HALF_LEAVE_CODES:
        return Decimal("0.50")
    return Decimal("0.00")


def _display_decimal(value: Decimal):
    """Trả số gọn cho template/Excel: 1 thay vì 1.00; 0.5 giữ 0.5."""
    value = _to_decimal(value)
    if value == value.to_integral_value():
        return int(value)
    return float(value.normalize())


def _employees_for_month(unit: OrgUnit, team_id: str | None, start: dt_date, end: dt_date):
    """
    Lấy danh sách nhân sự hiển thị cho báo cáo tháng:
    - Nhân sự ACTIVE thuộc đơn vị.
    - Nhân sự có dòng công chốt của đơn vị trong tháng.
    - Nhân sự có điều động đi/đến liên quan đến đơn vị trong tháng.

    TempAssignment là nguồn nghiệp vụ gốc cho BS. Khi lọc tổ, BS đi theo tổ
    hiện tại của nhân sự gốc; BS đến chỉ hiển thị khi xem toàn đơn vị vì chưa
    có trường "tổ tiếp nhận" tại đơn vị nhận.
    """
    base_qs = Employee.objects.filter(unit=unit, status=Employee.Status.ACTIVE)
    team_id_int = None
    if team_id:
        try:
            team_id_int = int(team_id)
            base_qs = base_qs.filter(team_id=team_id_int)
        except Exception:
            team_id_int = None

    base_ids = set(base_qs.values_list("id", flat=True))

    commits = AttendanceCommit.objects.filter(unit=unit, work_date__gte=start, work_date__lte=end).values_list("id", flat=True)
    items_qs = AttendanceCommitItem.objects.filter(commit_id__in=list(commits))
    if team_id_int:
        items_qs = items_qs.filter(employee__team_id=team_id_int)
    commit_emp_ids: Set[int] = set(items_qs.values_list("employee_id", flat=True))

    bs_emp_ids: Set[int] = set()
    ta_qs = _effective_temp_assignment_qs_between(start, end)
    if ta_qs is not None:
        # BS đi: người thuộc đơn vị gốc, lọc theo tổ được.
        bs_out_qs = ta_qs.filter(from_unit=unit)
        if team_id_int:
            bs_out_qs = bs_out_qs.filter(employee__team_id=team_id_int)
        bs_emp_ids.update(bs_out_qs.values_list("employee_id", flat=True))

        # BS đến: hiển thị khi xem toàn đơn vị. Khi lọc tổ thì chưa có tổ tiếp nhận để lọc chính xác.
        if not team_id_int:
            bs_emp_ids.update(ta_qs.filter(to_unit=unit).values_list("employee_id", flat=True))

    all_ids = base_ids.union(commit_emp_ids).union(bs_emp_ids)

    return Employee.objects.filter(id__in=list(all_ids)).order_by("employee_code")


def _commit_items_map(unit: OrgUnit, start: dt_date, end: dt_date) -> Dict[Tuple[int, dt_date], AttendanceCommitItem]:
    """Tải tất cả commit items theo đơn vị và khoảng ngày -> map (employee_id, work_date) => item."""
    commits = AttendanceCommit.objects.filter(unit=unit, work_date__gte=start, work_date__lte=end).values_list("id", "work_date")
    commit_ids = [cid for cid, _ in commits]
    work_date_by_commit = {cid: wd for cid, wd in commits}

    items = AttendanceCommitItem.objects.select_related("employee", "code", "bs_peer_unit").filter(commit_id__in=commit_ids)
    m: Dict[Tuple[int, dt_date], AttendanceCommitItem] = {}
    for it in items:
        wd = work_date_by_commit.get(it.commit_id)
        if wd:
            m[(it.employee_id, wd)] = it
    return m


def _bs_out_credit_map(unit: OrgUnit, start: dt_date, end: dt_date) -> Dict[Tuple[int, dt_date, int | None], Decimal]:
    """
    Với BS đi tại đơn vị gốc, công chính xác lấy từ dòng BS đến của đơn vị nhận.
    TempAssignment là nguồn xác định người/ngày/đơn vị nhận; commit bên nhận chỉ
    cung cấp số công thực tế nếu đã chốt.

    Map key: (employee_id, work_date, receiving_unit_id) -> work_credit của dòng nhận.
    """
    ta_qs = _effective_temp_assignment_qs_between(start, end)
    if ta_qs is None:
        return {}

    out_assignments = list(ta_qs.filter(from_unit=unit).select_related("to_unit", "employee"))
    if not out_assignments:
        return {}

    receiving_unit_ids = {ta.to_unit_id for ta in out_assignments if ta.to_unit_id}
    employee_ids = {ta.employee_id for ta in out_assignments}

    commits = AttendanceCommit.objects.filter(
        unit_id__in=receiving_unit_ids,
        work_date__gte=start,
        work_date__lte=end,
    ).values_list("id", "work_date", "unit_id")
    commit_meta = {cid: (wd, unit_id) for cid, wd, unit_id in commits}

    in_items = AttendanceCommitItem.objects.select_related("code").filter(
        commit_id__in=list(commit_meta.keys()),
        employee_id__in=employee_ids,
        bs_direction=AttendanceCommitItem.BSDirection.IN,
        include_in_unit=True,
        bs_peer_unit=unit,
    )

    actual_credit: Dict[Tuple[int, dt_date, int | None], Decimal] = {}
    for it in in_items:
        wd, receiving_unit_id = commit_meta.get(it.commit_id, (None, None))
        if not wd:
            continue
        actual_credit[(it.employee_id, wd, receiving_unit_id)] = _credit_from_item(it, "work_credit_snapshot", "work_credit")

    result: Dict[Tuple[int, dt_date, int | None], Decimal] = {}
    for ta in out_assignments:
        for d in _iter_effective_dates_for_assignment(ta, start, end):
            key = (ta.employee_id, d, ta.to_unit_id)
            if key in actual_credit:
                result[key] = actual_credit[key]
    return result


def _virtual_bs_out_item_map_from_temp_assignment(unit: OrgUnit, start: dt_date, end: dt_date) -> Dict[Tuple[int, dt_date], object]:
    """
    Tạo map dòng BS đi hiển thị cho đơn vị gốc từ điều động nhân sự.
    Không ghi database; chỉ dùng để báo cáo tháng không bị thiếu mã BS khi commit
    gốc chưa có dòng OUT.
    """
    ta_qs = _effective_temp_assignment_qs_between(start, end)
    if ta_qs is None:
        return {}

    result: Dict[Tuple[int, dt_date], object] = {}
    for ta in ta_qs.filter(from_unit=unit).order_by("employee__employee_code", "start_date", "id"):
        for d in _iter_effective_dates_for_assignment(ta, start, end):
            result[(ta.employee_id, d)] = SimpleNamespace(
                employee=ta.employee,
                employee_id=ta.employee_id,
                code=None,
                code_id=None,
                code_snapshot="BS",
                label_snapshot="Bổ sung đi",
                bs_direction=AttendanceCommitItem.BSDirection.OUT,
                bs_peer_unit=ta.to_unit,
                bs_peer_unit_id=ta.to_unit_id,
                include_in_unit=False,
                overtime_hours=0,
                in1=None,
                out1=None,
                in2=None,
                out2=None,
                notes=f"Bổ sung đi sang {ta.to_unit.code if ta.to_unit else ''}".strip(),
                is_virtual_bs_out=True,
                source_temp_assignment_id=ta.id,
            )
    return result


def _bs_monthly_warnings(unit: OrgUnit, start: dt_date, end: dt_date) -> List[str]:
    """Cảnh báo nhanh các trường hợp điều động và công chốt BS chưa khớp trong tháng."""
    warnings: List[str] = []
    ta_qs = _effective_temp_assignment_qs_between(start, end)
    if ta_qs is None:
        return warnings

    # Chuẩn bị commit item theo key để kiểm nhanh.
    related_unit_ids = {unit.id}
    for ta in ta_qs.filter(Q(from_unit=unit) | Q(to_unit=unit)):
        if ta.from_unit_id:
            related_unit_ids.add(ta.from_unit_id)
        if ta.to_unit_id:
            related_unit_ids.add(ta.to_unit_id)

    commits = AttendanceCommit.objects.filter(
        unit_id__in=related_unit_ids,
        work_date__gte=start,
        work_date__lte=end,
    ).values_list("id", "unit_id", "work_date")
    commit_meta = {cid: (unit_id, wd) for cid, unit_id, wd in commits}

    items = AttendanceCommitItem.objects.select_related("employee", "bs_peer_unit").filter(commit_id__in=list(commit_meta.keys()))
    item_by_unit_emp_date: Dict[Tuple[int, int, dt_date], AttendanceCommitItem] = {}
    for it in items:
        unit_id, wd = commit_meta.get(it.commit_id, (None, None))
        if unit_id and wd:
            item_by_unit_emp_date[(unit_id, it.employee_id, wd)] = it

    max_messages = 20
    hidden_count = 0

    def add_warning(msg: str):
        nonlocal hidden_count
        if len(warnings) < max_messages:
            warnings.append(msg)
        else:
            hidden_count += 1

    for ta in ta_qs.filter(from_unit=unit).order_by("employee__employee_code", "start_date", "id"):
        emp_label = f"{ta.employee.employee_code} - {ta.employee.full_name}"
        for d in _iter_effective_dates_for_assignment(ta, start, end):
            origin_item = item_by_unit_emp_date.get((unit.id, ta.employee_id, d))
            if origin_item:
                valid_out = (
                    str(origin_item.bs_direction) == str(AttendanceCommitItem.BSDirection.OUT)
                    and not bool(origin_item.include_in_unit)
                    and (origin_item.bs_peer_unit_id or None) == ta.to_unit_id
                )
                if not valid_out:
                    add_warning(f"{d.strftime('%d/%m')}: {emp_label} có điều động đi {ta.to_unit.code}, nhưng dòng chốt đơn vị gốc chưa phải BS đi đúng.")
            received_item = item_by_unit_emp_date.get((ta.to_unit_id, ta.employee_id, d))
            valid_received = bool(
                received_item
                and str(received_item.bs_direction) == str(AttendanceCommitItem.BSDirection.IN)
                and bool(received_item.include_in_unit)
                and (received_item.bs_peer_unit_id or None) == unit.id
            )
            if not valid_received:
                add_warning(f"{d.strftime('%d/%m')}: {emp_label} có điều động đi {ta.to_unit.code}, nhưng đơn vị nhận chưa có công chốt BS đến tương ứng.")

    for ta in ta_qs.filter(to_unit=unit).order_by("employee__employee_code", "start_date", "id"):
        emp_label = f"{ta.employee.employee_code} - {ta.employee.full_name}"
        for d in _iter_effective_dates_for_assignment(ta, start, end):
            received_item = item_by_unit_emp_date.get((unit.id, ta.employee_id, d))
            valid_in = bool(
                received_item
                and str(received_item.bs_direction) == str(AttendanceCommitItem.BSDirection.IN)
                and bool(received_item.include_in_unit)
                and (received_item.bs_peer_unit_id or None) == ta.from_unit_id
            )
            if not valid_in:
                add_warning(f"{d.strftime('%d/%m')}: {emp_label} có điều động đến từ {ta.from_unit.code}, nhưng công chốt đơn vị nhận chưa phải BS đến đúng.")

    if hidden_count:
        warnings.append(f"Còn {hidden_count} cảnh báo BS khác chưa hiển thị. Vui lòng xuất/kiểm tra theo đơn vị-ngày nếu cần.")
    return warnings


def _build_monthly_rows(employees, days, items_map, bs_out_map):
    rows: List[Dict] = []
    for idx, emp in enumerate(employees, start=1):
        cells: List[Dict] = []
        # T.Làm: công làm thực tế tại đơn vị đang xem, gồm cả người BS đến.
        # T.BS đi: công thực tế của người thuộc đơn vị đi làm ở đơn vị nhận.
        # T.Thưởng: theo quy tắc hiện tại = T.Làm + T.BS đi.
        total_work = Decimal("0.00")
        total_overtime = 0
        total_weekend_work = Decimal("0.00")
        total_leave = Decimal("0.00")
        total_bs_in = Decimal("0.00")
        total_bs_out = Decimal("0.00")

        for d in days:
            it = items_map.get((emp.id, d))
            code_text = _commit_item_code_text(it)
            bs_direction = getattr(it, "bs_direction", "") if it else ""
            peer_unit = getattr(it, "bs_peer_unit", None) if it else None

            work_credit = Decimal("0.00")
            overtime_hours = 0

            if it and bool(getattr(it, "include_in_unit", True)):
                work_credit = _credit_from_item(it, "work_credit_snapshot", "work_credit")
                overtime_hours = int(getattr(it, "overtime_hours", 0) or 0)
                total_work += work_credit
                total_overtime += overtime_hours
                total_leave += _leave_credit(it)
                if d.weekday() >= 5:
                    total_weekend_work += work_credit
                if _is_supplement_in(it):
                    total_bs_in += work_credit
            elif it and _is_supplement_out(it):
                peer_unit_id = getattr(it, "bs_peer_unit_id", None)
                total_bs_out += bs_out_map.get((emp.id, d, peer_unit_id), Decimal("0.00"))

            cells.append({
                "code": code_text,
                "bs_direction": bs_direction,
                "peer_unit": peer_unit,
                "is_weekend": d.weekday() >= 5,
            })

        total_reward = total_work + total_bs_out

        rows.append({
            "idx": idx,
            "emp": emp,
            "cells": cells,
            "total_work": _display_decimal(total_work),
            "total_bs_out": _display_decimal(total_bs_out),
            "total_reward": _display_decimal(total_reward),
            "total_overtime": total_overtime,
            "total_weekend_work": _display_decimal(total_weekend_work),
            "total_leave": _display_decimal(total_leave),
            "total_bs_in": _display_decimal(total_bs_in),
        })
    return rows


def _build_monthly_dataset(unit: OrgUnit, team_id: str | None, start: dt_date, end: dt_date, days: List[dt_date]) -> Dict:
    """Dữ liệu dùng chung cho màn hình và Excel để tránh lệch cách tính."""
    employees = _employees_for_month(unit, team_id, start, end)
    items_map = _commit_items_map(unit, start, end)
    for key, virtual_item in _virtual_bs_out_item_map_from_temp_assignment(unit, start, end).items():
        items_map.setdefault(key, virtual_item)
    bs_out_map = _bs_out_credit_map(unit, start, end)
    rows = _build_monthly_rows(employees, days, items_map, bs_out_map)
    bs_warnings = _bs_monthly_warnings(unit, start, end)
    return {
        "employees": employees,
        "items_map": items_map,
        "bs_out_map": bs_out_map,
        "rows": rows,
        "bs_warnings": bs_warnings,
    }


@login_required
def monthly_view(request):
    """
    Xem chấm công theo tháng (từ dữ liệu đã chốt).
    Bộ lọc: đơn vị, tổ, tháng, năm.
    """
    if not request.user.has_perm("attendance.view_attendancecommit"):
        return render(request, "backoffice/no_permission.html", {
            "perm_codename": "attendance.view_attendancecommit",
            "title": "Bạn chưa được cấp quyền xem Công theo tháng"
        }, status=403)

    units_qs = _units_for_attendance(request)

    unit_id = (request.GET.get("unit") or "").strip()
    team_id = (request.GET.get("team") or "").strip()
    month, year = _parse_month_year(request)

    unit = None
    if unit_id:
        try:
            unit = OrgUnit.objects.get(pk=int(unit_id), is_attendance_unit=True, is_active=True)
        except Exception:
            unit = None
            messages.warning(request, "Đơn vị không hợp lệ hoặc không được phép chấm công.")
        if unit and not unit_in_attendance_scope(request.user, unit.id):
            unit = None
            messages.error(request, "Bạn không có quyền xem đơn vị này.")

    teams = OrgUnit.objects.filter(parent=unit, type=OrgUnit.Type.TEAM).order_by("symbol") if unit else OrgUnit.objects.none()

    start, end, dim = _month_range(year, month)
    days = [start + timedelta(days=i) for i in range(dim)]

    dataset = _build_monthly_dataset(unit, team_id, start, end, days) if unit else {"rows": [], "bs_warnings": []}
    rows = dataset["rows"]
    bs_warnings = dataset["bs_warnings"]

    return render(request, "backoffice/attendance/monthly_view.html", {
        "units_qs": units_qs,
        "unit_id": unit_id,
        "month": month,
        "year": year,
        "teams": teams,
        "team_id": team_id,
        "days": [{"d": d.day, "wd": _weekday_label_precise(d), "is_weekend": d.weekday() >= 5} for d in days],
        "rows": rows,
        "bs_warnings": bs_warnings,
        "has_data": bool(unit),
    })


@login_required
def monthly_export(request):
    """
    Xuất Excel: BẢNG CHẤM CÔNG CỦA "ĐƠN VỊ".
    Dữ liệu xuất đồng bộ với bảng xem tháng, gồm các cột tổng và BS đi/BS đến.
    """
    if not request.user.has_perm("attendance.view_attendancecommit"):
        return render(request, "backoffice/no_permission.html", {
            "perm_codename": "attendance.view_attendancecommit",
            "title": "Bạn chưa được cấp quyền xuất Công theo tháng"
        }, status=403)

    if openpyxl is None:
        messages.error(request, "Thiếu thư viện openpyxl để xuất Excel.")
        return render(request, "backoffice/attendance/monthly_view.html", {"has_data": False})

    unit_id = (request.GET.get("unit") or "").strip()
    team_id = (request.GET.get("team") or "").strip()
    month, year = _parse_month_year(request)

    try:
        unit = OrgUnit.objects.get(pk=int(unit_id), is_attendance_unit=True, is_active=True)
    except Exception:
        messages.error(request, "Đơn vị không hợp lệ.")
        return render(request, "backoffice/attendance/monthly_view.html", {"has_data": False})
    if not unit_in_attendance_scope(request.user, unit.id):
        messages.error(request, "Bạn không có quyền xuất dữ liệu đơn vị này.")
        return render(request, "backoffice/attendance/monthly_view.html", {"has_data": False})

    start, end, dim = _month_range(year, month)
    days = [start + timedelta(days=i) for i in range(dim)]

    dataset = _build_monthly_dataset(unit, team_id, start, end, days)
    rows = dataset["rows"]
    bs_warnings = dataset["bs_warnings"]

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Bang cham cong"

    thin = Side(style="thin", color="000000")
    border = Border(top=thin, left=thin, right=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    weekend_fill = PatternFill("solid", fgColor="FFF3CD")
    total_fill = PatternFill("solid", fgColor="E7F1FF")

    total_headers = ["T.Làm", "T.BS đi", "T.Thưởng", "T.Thêm", "T.T7CN", "T.Nghỉ", "T.BS đến"]
    total_cols = 2 + dim + len(total_headers)

    title_text = f'BẢNG CHẤM CÔNG CỦA "{unit.name}"'
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    ws.cell(row=1, column=1, value=title_text).font = Font(bold=True, size=14)
    ws.cell(row=1, column=1).alignment = center

    ws.cell(row=2, column=1, value=f"Tháng {month:02d}/{year}").font = Font(bold=True)
    ws.cell(row=2, column=1).alignment = Alignment(horizontal="left")
    ws.cell(row=2, column=2, value="Họ và tên").font = Font(bold=True)
    ws.cell(row=2, column=2).alignment = center

    for i, d in enumerate(days, start=1):
        col = 2 + i
        cell = ws.cell(row=2, column=col, value=f"{d.day:02d}\n{_weekday_label_precise(d)}")
        cell.alignment = center
        cell.font = Font(bold=True)
        if d.weekday() >= 5:
            cell.fill = weekend_fill

    first_total_col = 3 + dim
    for offset, header in enumerate(total_headers):
        col = first_total_col + offset
        cell = ws.cell(row=2, column=col, value=header)
        cell.alignment = center
        cell.font = Font(bold=True)
        cell.fill = total_fill

    row_idx = 3
    for row in rows:
        ws.cell(row=row_idx, column=1, value=row["idx"]).alignment = center
        ws.cell(row=row_idx, column=2, value=f"{row['emp'].full_name}").alignment = Alignment(horizontal="left")

        for i, cell_data in enumerate(row["cells"], start=1):
            col = 2 + i
            code_value = cell_data.get("code") or ""
            bs_direction = cell_data.get("bs_direction") or ""
            if bs_direction == AttendanceCommitItem.BSDirection.IN and code_value:
                code_value = f"{code_value}+"
            c = ws.cell(row=row_idx, column=col, value=code_value)
            c.alignment = center
            if cell_data.get("is_weekend"):
                c.fill = weekend_fill

        totals = [
            row["total_work"], row["total_bs_out"], row["total_reward"], row["total_overtime"],
            row["total_weekend_work"], row["total_leave"], row["total_bs_in"],
        ]
        for offset, value in enumerate(totals):
            c = ws.cell(row=row_idx, column=first_total_col + offset, value=value)
            c.alignment = center
            c.fill = total_fill
        row_idx += 1

    for r in range(1, row_idx):
        for c in range(1, total_cols + 1):
            ws.cell(row=r, column=c).border = border

    ws.column_dimensions[get_column_letter(1)].width = 6
    ws.column_dimensions[get_column_letter(2)].width = 28
    for i in range(dim):
        ws.column_dimensions[get_column_letter(3 + i)].width = 5
    for c in range(first_total_col, total_cols + 1):
        ws.column_dimensions[get_column_letter(c)].width = 10

    ws.freeze_panes = ws["C3"]

    if bs_warnings:
        ws_warn = wb.create_sheet("Canh bao BS")
        ws_warn.cell(row=1, column=1, value="Cảnh báo BS đi/BS đến theo điều động").font = Font(bold=True, size=13)
        ws_warn.cell(row=2, column=1, value="Các cảnh báo này không làm thay đổi dữ liệu đã chốt; dùng để thống kê kiểm tra lại đơn vị/ngày liên quan.")
        for idx, warning in enumerate(bs_warnings, start=4):
            ws_warn.cell(row=idx, column=1, value=warning)
            ws_warn.cell(row=idx, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws_warn.column_dimensions[get_column_letter(1)].width = 120

    from django.utils.encoding import iri_to_uri
    filename = f"bang-cham-cong-{unit.symbol or unit.code}-{year}-{month:02d}.xlsx"
    response = HttpResponse(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{iri_to_uri(filename)}"'
    wb.save(response)
    return response
