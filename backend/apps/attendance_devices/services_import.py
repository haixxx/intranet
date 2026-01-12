import csv
from io import TextIOWrapper, BytesIO
from typing import List, Dict, Tuple, Optional
from datetime import datetime, date, time

from django.db import transaction
from django.utils import timezone

from apps.hr.models import Employee
from .models_adjustments import AttendanceImportBatch, AttendanceImportLine
from .services_overlay import overlay_day


def parse_time_str(s: str) -> Optional[time]:
    s = (s or "").strip()
    if not s:
        return None
    try:
        hh, mm = [int(x) for x in s.split(":")]
        return time(hour=hh, minute=mm)
    except Exception:
        return None


def parse_upload(file) -> Tuple[List[Dict], List[str]]:
    """
    Nhận cả CSV và XLSX:
    - CSV header: employee_code,work_date,in1,out1,in2,out2,note
    - XLSX: sheet đầu, dòng 1 là header tương tự CSV
    Trả về (rows, errors)
    """
    name = (getattr(file, "name", "") or "").lower()
    rows: List[Dict] = []
    errors: List[str] = []

    if name.endswith(".csv"):
        # csv utf-8
        wrapper = TextIOWrapper(file, encoding="utf-8")
        reader = csv.DictReader(wrapper)
        required = {"employee_code", "work_date"}
        if not required.issubset(set(reader.fieldnames or [])):
            errors.append("CSV phải có cột employee_code, work_date.")
            return [], errors
        for idx, r in enumerate(reader, start=2):  # line 2 là dữ liệu
            row = {
                "employee_code": (r.get("employee_code") or "").strip(),
                "work_date": (r.get("work_date") or "").strip(),
                "in1": parse_time_str(r.get("in1")),
                "out1": parse_time_str(r.get("out1")),
                "in2": parse_time_str(r.get("in2")),
                "out2": parse_time_str(r.get("out2")),
                "note": (r.get("note") or "").strip(),
                "_line": idx,
            }
            rows.append(row)
        return rows, errors
    elif name.endswith(".xlsx"):
        try:
            import openpyxl  # yêu cầu cài đặt openpyxl
        except ImportError:
            errors.append("Thiếu thư viện openpyxl để đọc XLSX. Vui lòng cài đặt: pip install openpyxl")
            return [], errors

        wb = openpyxl.load_workbook(BytesIO(file.read()), data_only=True)
        ws = wb.active
        headers = [str((ws.cell(row=1, column=i).value or "")).strip() for i in range(1, ws.max_column + 1)]
        expected = ["employee_code", "work_date", "in1", "out1", "in2", "out2", "note"]
        if not set(["employee_code", "work_date"]).issubset(set(headers)):
            errors.append("XLSX phải có cột employee_code, work_date.")
            return [], errors

        idx_map = {h: headers.index(h) + 1 for h in headers}
        for r in range(2, ws.max_row + 1):
            def cell(h):
                col = idx_map.get(h)
                return ws.cell(row=r, column=col).value if col else None

            row = {
                "employee_code": str(cell("employee_code") or "").strip(),
                "work_date": str(cell("work_date") or "").strip(),
                "in1": parse_time_str(str(cell("in1") or "")),
                "out1": parse_time_str(str(cell("out1") or "")),
                "in2": parse_time_str(str(cell("in2") or "")),
                "out2": parse_time_str(str(cell("out2") or "")),
                "note": str(cell("note") or "").strip(),
                "_line": r,
            }
            rows.append(row)
        return rows, errors
    else:
        errors.append("Định dạng tệp không hỗ trợ. Chỉ nhận .csv hoặc .xlsx")
        return [], errors


@transaction.atomic
def apply_import(rows: List[Dict], user, mode: str, reason_code: str, note: str) -> AttendanceImportBatch:
    """
    Tạo batch + lines từ rows hợp lệ, sau đó đánh dấu applied và overlay từng ngày.
    """
    # batch_id sinh từ thời gian và user
    batch_id = f"DFI-{user.id}-{timezone.now().strftime('%Y%m%d%H%M%S')}"
    batch = AttendanceImportBatch.objects.create(
        batch_id=batch_id,
        file_name="upload",
        imported_by=user,
        mode=mode,
        reason_code=reason_code or "",
        note=note or "",
    )

    applied = 0
    errors = 0
    for row in rows:
        code = row["employee_code"]
        wd_str = row["work_date"]
        try:
            y, m, d = [int(x) for x in wd_str.split("-")]
            wd = date(y, m, d)
        except Exception:
            line = AttendanceImportLine.objects.create(
                batch=batch, employee=None, work_date=date(1970, 1, 1),
                valid=False, error_msg=f"Dòng {row.get('_line')}: work_date không đúng định dạng YYYY-MM-DD",
            )
            errors += 1
            continue

        emp = Employee.objects.filter(employee_code=code).first()
        if not emp:
            AttendanceImportLine.objects.create(
                batch=batch, employee=None, work_date=wd,
                valid=False, error_msg=f"Dòng {row.get('_line')}: Không tìm thấy nhân sự {code}",
            )
            errors += 1
            continue

        line = AttendanceImportLine.objects.create(
            batch=batch,
            employee=emp,
            work_date=wd,
            in1=row.get("in1"),
            out1=row.get("out1"),
            in2=row.get("in2"),
            out2=row.get("out2"),
            reason_code=reason_code or "",
            note=row.get("note") or "",
            valid=True,
            applied=False,
        )
        # Áp overlay ngay cho ngày đó
        overlay_day(emp.id, wd)
        line.applied = True
        line.save(update_fields=["applied"])
        applied += 1

    batch.summary_json = {"applied": applied, "errors": errors}
    batch.save(update_fields=["summary_json"])
    return batch